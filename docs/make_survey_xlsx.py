from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
wb = Workbook()
ws = wb.active
ws.title = 'Survey Questionnaire'
ws['A1'] = 'PATHPOINT: A Role-Based Job Application and Interview Management Platform'
ws['A1'].font = Font(bold=True, size=14)
ws.merge_cells('A1:J1')
ws['A2'] = 'SURVEY QUESTIONNAIRE'
ws['A2'].font = Font(bold=True, size=12)
ws.merge_cells('A2:J2')
ws['A4'] = 'Name (Optional):'
ws['B4'] = '___________________________'
ws['I4'] = 'Date:'
ws['J4'] = '___________________________'
ws['A6'] = 'Direction: Please indicate your level of agreement with each statement by placing a check mark (/) in the appropriate column.'
ws.merge_cells('A6:J6')
ws['A7'] = 'Scale: 1 = Strongly Disagree (SD) | 2 = Disagree (D) | 3 = Neutral (N) | 4 = Agree (A) | 5 = Strongly Agree (SA)'
ws.merge_cells('A7:J7')
ws['A9'] = 'A. RESPONDENT PROFILE'
ws['A9'].font = Font(bold=True)
ws['A10'] = '1. Age:'
ws['B10'] = '_____'
ws['A11'] = '2. Sex:'
ws['B11'] = u'\u2610 Male   \u2610 Female   \u2610 LGBTQIA+   \u2610 Prefer not to say'
ws['A12'] = '3. Role:'
ws['B12'] = u'\u2610 Job Seeker/Applicant   \u2610 Employer/Provider   \u2610 Career Counselor/School Personnel'
ws['A13'] = '4. Location (City/Municipality):'
ws['B13'] = '___________________________'
ws.merge_cells('B11:J11')
ws.merge_cells('B12:J12')
ws.merge_cells('B13:J13')
thin = Side(style='thin')
border = Border(left=thin,right=thin,top=thin,bottom=thin)
fill = PatternFill('solid', fgColor='D9D9D9')

def add_section(start_row, title, rows):
    r = start_row
    ws.cell(row=r, column=1, value=title).font = Font(bold=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    r += 1
    headers = ['NO.', 'STATEMENT', 'SD (1)', 'D (2)', 'N (3)', 'A (4)', 'SA (5)']
    for i,h in enumerate(headers,1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.fill = fill
        c.border = border
    r += 1
    for no,stmt in rows:
        ws.cell(row=r,column=1,value=no).alignment = Alignment(horizontal='center')
        ws.cell(row=r,column=2,value=stmt).alignment = Alignment(wrap_text=True, vertical='top')
        for c in range(3,8):
            cell = ws.cell(row=r,column=c,value=u'\u2610')
            cell.alignment = Alignment(horizontal='center')
        for c in range(1,8):
            ws.cell(row=r,column=c).border = border
        ws.row_dimensions[r].height = 36
        r += 1
    return r

r = 16
r = add_section(r, 'B. PERCEIVED PROBLEMS IN JOB APPLICATION AND HIRING PROCESS (For All Respondents)', [
    (1,'It is difficult to keep track of job applications and their statuses.'),
    (2,'Important deadlines (application deadlines/interviews) are often missed due to poor tracking.'),
    (3,'Managing application requirements (resume, documents, etc.) is time-consuming and hassle.'),
    (4,'Communication between applicants and employers regarding job applications is not efficient.'),
    (5,'There is a need for a centralized platform to manage job applications and hiring processes.'),
])
r += 1
r = add_section(r, 'C. PERCEIVED USEFULNESS OF PATHPOINT FOR JOB SEEKERS/APPLICANTS', [
    (6,'PathPoint can help me easily track all my job applications in one place.'),
    (7,'The application tracker (Saved/Applied/Interviewing/Offered/Rejected) is useful for monitoring my progress.'),
    (8,'Having a list of application requirements for each job posting will help me prepare better.'),
    (9,'The work log feature can help me document my skills and work experiences.'),
    (10,'PathPoint can help me stay organized and avoid missing interviews or deadlines.'),
])
r += 1
r = add_section(r, 'D. PERCEIVED USEFULNESS OF PATHPOINT FOR EMPLOYERS/PROVIDERS', [
    (11,'PathPoint can help employers post and manage job openings more efficiently.'),
    (12,'The system can help in screening and reviewing applicants based on their submitted requirements.'),
    (13,'Updating applicant statuses (Shortlisted, Interview Scheduled, Accepted, Rejected) is convenient through the system.'),
    (14,'The platform can reduce the use of manual methods (paper forms, Messenger, etc.) in hiring.'),
    (15,'Company verification adds credibility and trust to job postings.'),
])
r += 1
r = add_section(r, 'E. PERCEVED USEFULNESS OF PATHPOINT FOR CAREER COUNSELORS/SCHOOL PERSONNEL', [
    (16,'PathPoint can help career counselors monitor their students/graduates job applications more effectively.'),
    (17,'The system can aid in providing better guidance and follow-up to job-seeking students.'),
    (18,'Having a centralized platform makes it easier to assist multiple students/graduates.'),
])
r += 1
r = add_section(r, 'F. OVERALL EVALUATION OF PATHPOINT', [
    (19,'PathPoint is relevant to the needs of job seekers and employers in the community.'),
    (20,'PathPoint is user-friendly and easy to navigate.'),
    (21,'PathPoint is a practical solution to the problems in job application and hiring processes.'),
    (22,'I am willing to use PathPoint if it is made available.'),
    (23,'I would recommend PathPoint to other job seekers, employers, or career counselors.'),
])
r += 2
ws.cell(row=r,column=1,value='G. SUGGESTIONS (Optional)').font = Font(bold=True)
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
r += 1
ws.cell(row=r,column=1,value='1. What features would you like to see added to PathPoint?')
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
r += 1
for i in range(3):
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
    ws.row_dimensions[r].height=20
    r += 1
ws.cell(row=r,column=1,value='2. Other comments and suggestions:')
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
r += 1
for i in range(3):
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
    ws.row_dimensions[r].height=20
    r += 1
ws.column_dimensions['A'].width=6
ws.column_dimensions['B'].width=70
for c in 'CDEFG':
    ws.column_dimensions[c].width=10
wb.save(r'C:\PathPoint\docs\SURVEY_QUESTIONNAIRE.xlsx')
print('saved')
